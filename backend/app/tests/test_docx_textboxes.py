from docx import Document
from docx.oxml import parse_xml

from app.tools.document_parser import parse_document


def test_textbox_compatibility_branch_is_read_once(tmp_path):
    document = Document()
    document.element.body.insert(
        0,
        parse_xml("""
      <w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
           xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">
        <w:r><mc:AlternateContent>
          <mc:Choice Requires="wps">
            <w:txbxContent><w:p><w:r><w:t>Python 项目经历</w:t></w:r></w:p></w:txbxContent>
          </mc:Choice>
          <mc:Fallback>
            <w:txbxContent><w:p><w:r><w:t>Python 项目经历</w:t></w:r></w:p></w:txbxContent>
          </mc:Fallback>
        </mc:AlternateContent></w:r>
      </w:p>"""),
    )
    path = tmp_path / "textbox.docx"
    document.save(path)
    assert parse_document(path, "")[0].content == "Python 项目经历"


def test_nested_table_hyperlink_and_deleted_text(tmp_path):
    document = Document()
    document.add_paragraph("开始")
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "嵌套经历"
    paragraph = document.add_paragraph()
    paragraph._p.append(
        parse_xml("""
      <w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
        <w:r><w:t>作品链接说明</w:t></w:r>
      </w:hyperlink>""")
    )
    paragraph._p.append(
        parse_xml("""
      <w:del xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
        <w:r><w:delText>已经删除的经历</w:delText></w:r>
      </w:del>""")
    )
    path = tmp_path / "nested.docx"
    document.save(path)
    assert parse_document(path, "")[0].content == "开始\n嵌套经历\n作品链接说明"
