import asyncio
import subprocess
import time
from datetime import UTC, datetime, timedelta
from types import SimpleNamespace
from uuid import UUID, uuid4

import fitz
import pytest
from docx import Document
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError

from app.core.exceptions import AppError
from app.models.analysis import AnalysisTask
from app.models.enums import TaskStatus
from app.models.identity import User
from app.services import portfolio, task_dispatcher, task_lifecycle
from app.services.task_lifecycle import (
    claim_task,
    finish_task,
    heartbeat_task,
    recover_interrupted_tasks,
    reset_failed_task,
)
from app.services.task_runner import run_isolated
from app.tests.test_portfolio_api import app_database, student  # noqa: F401
from app.tools.document_parser import parse_document


def test_docx_tables_are_extracted_and_blank_pdf_without_ocr_fails(tmp_path, monkeypatch):
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "经历"
    table.cell(0, 1).text = "用户调研"
    table.cell(1, 0).text = "成果"
    table.cell(1, 1).text = "20份访谈"
    path = tmp_path / "table.docx"
    document.save(path)
    text = parse_document(path, "")[0].content
    assert "用户调研" in text and "20份访谈" in text
    pdf = tmp_path / "blank.pdf"
    with fitz.open() as empty:
        empty.new_page()
        empty.save(pdf)
    monkeypatch.setattr("app.tools.document_parser.settings.ocr_enabled", False)
    with pytest.raises(AppError, match="未提取或识别"):
        parse_document(pdf, "")


async def test_isolated_timeout_kills_process_and_keeps_event_loop_responsive(monkeypatch):
    processes = []
    original = subprocess.Popen

    def spawn(*args, **kwargs):
        process = original(*args, **kwargs)
        processes.append(process)
        return process

    monkeypatch.setattr(subprocess, "Popen", spawn)
    ticked = asyncio.Event()

    async def tick():
        await asyncio.sleep(0.05)
        ticked.set()

    ticker = asyncio.create_task(tick())
    with pytest.raises(AppError, match="超时"):
        await run_isolated(time.sleep, 10, timeout_seconds=0.3)
    await ticker
    assert ticked.is_set()
    assert processes and processes[0].returncode is not None


async def test_task_claim_retry_and_late_result_are_atomic(student):  # noqa: F811
    async with portfolio.async_session_factory() as session:
        user = await session.scalar(select(User).where(User.email == "alice@example.com"))
        task = AnalysisTask(created_by=user.id, task_type="test", status=TaskStatus.PENDING)
        session.add(task)
        await session.commit()
        task_id = task.id
        first_token = await claim_task(session, task_id, step="first")
        assert first_token is not None
        assert await claim_task(session, task_id, step="duplicate") is None
        assert not await finish_task(session, task_id, uuid4(), step="wrong token")
        await session.rollback()
        assert await finish_task(
            session, task_id, first_token, status=TaskStatus.FAILED, step="failed"
        )
        await session.commit()
        await session.refresh(user)
        await reset_failed_task(session, task_id, user)
        await session.commit()
        second_token = await claim_task(session, task_id, step="retry")
        assert second_token is not None and second_token != first_token
        assert not await finish_task(session, task_id, first_token, step="late result")
        await session.rollback()
        assert await finish_task(session, task_id, second_token, step="done")
        await session.commit()
        task = await session.get(AnalysisTask, task_id, populate_existing=True)
        assert task.status == TaskStatus.COMPLETED and task.attempt == 2


async def test_recovery_expires_stale_pending_and_running_but_not_live_tasks(student):  # noqa: F811
    now = datetime.now(UTC)
    old = now - timedelta(hours=1)
    async with portfolio.async_session_factory() as session:
        user = await session.scalar(select(User).where(User.email == "alice@example.com"))
        stale_running = AnalysisTask(
            created_by=user.id, task_type="test", status=TaskStatus.RUNNING, heartbeat_at=old
        )
        stale_pending = AnalysisTask(
            created_by=user.id, task_type="test", status=TaskStatus.PENDING, updated_at=old
        )
        live = AnalysisTask(
            created_by=user.id, task_type="test", status=TaskStatus.RUNNING, heartbeat_at=now
        )
        session.add_all([stale_running, stale_pending, live])
        await session.commit()
        ids = [task.id for task in (stale_running, stale_pending, live)]
    assert await recover_interrupted_tasks(portfolio.async_session_factory) == 2
    async with portfolio.async_session_factory() as session:
        statuses = [(await session.get(AnalysisTask, task_id)).status for task_id in ids]
        assert statuses == [TaskStatus.FAILED, TaskStatus.FAILED, TaskStatus.RUNNING]


@pytest.mark.parametrize("failure_stage", ["execute", "commit"])
async def test_heartbeat_retries_after_transient_database_failure(monkeypatch, failure_stage):
    attempts = 0
    commits = 0

    class Session:
        async def __aenter__(self):
            return self

        async def __aexit__(self, _type, _value, _traceback):
            return None

        async def execute(self, _statement):
            nonlocal attempts
            attempts += 1
            if attempts == 1 and failure_stage == "execute":
                raise SQLAlchemyError("temporary database outage")
            return SimpleNamespace(rowcount=1 if attempts <= 2 else 0)

        async def commit(self):
            nonlocal commits
            commits += 1
            if commits == 1 and failure_stage == "commit":
                raise SQLAlchemyError("temporary database outage")

    monkeypatch.setattr(task_lifecycle.settings, "task_lease_seconds", 0)
    await heartbeat_task(uuid4(), uuid4(), Session)
    assert attempts == 3
    assert commits == (2 if failure_stage == "execute" else 3)


async def test_failed_material_retry_preserves_version_and_replay(student):  # noqa: F811
    response = await student.post(
        "/api/v1/resumes",
        data={"title": "empty"},
        files={"file": ("empty.txt", b" ", "text/plain")},
    )
    accepted = response.json()["data"]
    task_id = UUID(accepted["taskId"])
    async with portfolio.async_session_factory() as session:
        user = await session.scalar(select(User).where(User.email == "alice@example.com"))
        retry, function, args = await portfolio.retry_material_task(session, task_id, user)
        assert retry.task_id == task_id
        assert str(args[1]) == accepted["versionId"]
    await function(*args)
    history = (await student.get("/api/v1/resumes")).json()["data"]
    assert len(history) == 1 and len(history[0]["versions"]) == 1
    assert history[0]["versions"][0]["status"] == "failed"


async def test_dispatcher_resumes_database_backed_pending_task(student, monkeypatch):  # noqa: F811
    async with portfolio.async_session_factory() as session:
        user = await session.scalar(select(User).where(User.email == "alice@example.com"))
        task = AnalysisTask(created_by=user.id, task_type="job_match", status=TaskStatus.PENDING)
        session.add(task)
        await session.commit()
        task_id = task.id

    async def complete(pending_id):
        async with portfolio.async_session_factory() as session:
            token = await claim_task(session, pending_id, step="自动恢复")
            assert token
            assert await finish_task(session, pending_id, token, step="恢复完成")
            await session.commit()

    monkeypatch.setattr(task_dispatcher, "execute_match_task", complete)
    assert await task_dispatcher.dispatch_pending_tasks(portfolio.async_session_factory) == 1
    await task_dispatcher.wait_for_dispatched_tasks()
    async with portfolio.async_session_factory() as session:
        recovered = await session.get(AnalysisTask, task_id)
        assert recovered.status == TaskStatus.COMPLETED
